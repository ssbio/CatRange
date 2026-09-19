#!/usr/bin/env python3
"""Read every workbook/CSV and optionally verify a previously reviewed manifest.

Usage:
    python tools/audit_tabular_files.py --json-output audit.json
    python tools/audit_tabular_files.py --baseline audit.json

Requires openpyxl for XLSX/XLSM. Does not modify input data. Parsing validates
file integrity, not scientific correctness or the completeness of an unknown
original. Comparing a reviewed hash manifest detects subsequent file changes,
including truncation that happens to end on a complete CSV record.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import zipfile


SUFFIXES = {".csv", ".xlsx", ".xlsm", ".xls", ".xlsb", ".ods"}
SKIP_DIRECTORIES = {".git", ".venv", "venv", "node_modules", "__pycache__", ".pytest_cache"}


def inventory(root: Path) -> list[Path]:
    found = []
    for directory, folders, files in os.walk(root, followlinks=False):
        folders[:] = sorted(name for name in folders if name not in SKIP_DIRECTORIES)
        for name in sorted(files):
            path = Path(directory) / name
            if path.suffix.lower() in SUFFIXES:
                found.append(path)
    return sorted(found)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def audit_csv(path: Path) -> dict:
    raw = path.read_bytes()
    if not raw:
        raise ValueError("Empty file: no CSV header")
    encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    text = raw.decode(encoding, errors="strict")
    if "\x00" in text:
        raise ValueError("NUL byte in decoded CSV")
    csv.field_size_limit(2**31 - 1)
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    try:
        header = next(reader)
    except StopIteration as error:
        raise ValueError("Missing CSV header") from error
    if not header or not any(value.strip() for value in header):
        raise ValueError("Empty CSV header")
    count = 0
    blank_records = 0
    mismatches = []
    for record_number, row in enumerate(reader, start=2):
        if not row:
            blank_records += 1
            continue
        count += 1
        if len(row) != len(header) and len(mismatches) < 10:
            mismatches.append({"record": record_number, "fields": len(row), "expected": len(header)})
    if mismatches:
        raise ValueError("Inconsistent CSV row widths: " + json.dumps(mismatches))
    warnings = []
    if not text.endswith(("\n", "\r")):
        warnings.append("No final newline; valid CSV, but check source completeness")
    if len(raw) >= 1024 * 1024 and len(raw) % (64 * 1024) == 0:
        warnings.append("Exact 64-KiB size boundary; compare with original for possible transfer truncation")
    return {"encoding": encoding, "data_rows": count, "columns": len(header),
            "header_only": count == 0, "blank_records": blank_records, "warnings": warnings}


def audit_workbook(path: Path) -> dict:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Unsupported workbook format: an appropriate reader is required")
    with zipfile.ZipFile(path) as archive:
        damaged_member = archive.testzip()
        if damaged_member is not None:
            raise ValueError(f"ZIP integrity error in {damaged_member}")
        for required in ("[Content_Types].xml", "xl/workbook.xml"):
            if required not in archive.namelist():
                raise ValueError(f"Missing workbook member: {required}")
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            declared_rows, declared_columns = sheet.max_row, sheet.max_column
            # Some writers store an incorrect dimension. Read through the
            # worksheet XML instead of stopping at a possibly truncated range.
            sheet.reset_dimensions()
            rows_read = 0
            for _ in sheet.iter_rows(values_only=True):
                rows_read += 1
            sheets.append({"name": sheet.title, "rows_including_header": rows_read,
                           "declared_rows": declared_rows, "declared_columns": declared_columns})
        if not sheets:
            raise ValueError("Workbook has no worksheets")
        return {"sheets": sheets, "warnings": []}
    finally:
        workbook.close()


def audit_file(path: Path, root: Path) -> dict:
    result = {"path": path.relative_to(root).as_posix(), "kind": path.suffix.lower()[1:]}
    try:
        before = path.stat()
        result.update(bytes=before.st_size, sha256=sha256(path))
        details = audit_csv(path) if path.suffix.lower() == ".csv" else audit_workbook(path)
        result.update(details)
        after = path.stat()
        if (before.st_size, before.st_mtime_ns, before.st_ino) != (after.st_size, after.st_mtime_ns, after.st_ino):
            raise ValueError("File changed during the audit")
        result["status"] = "pass"
    except Exception as error:
        result.update(status="fail", error=f"{type(error).__name__}: {error}")
    return result


def compare_baseline(results: list[dict], baseline: dict) -> list[dict]:
    expected = {item["path"]: item for item in baseline["files"]}
    actual = {item["path"]: item for item in results}
    changes = []
    for path in sorted(expected.keys() | actual.keys()):
        if path not in actual:
            changes.append({"path": path, "change": "missing"})
        elif path not in expected:
            changes.append({"path": path, "change": "added"})
        elif actual[path].get("sha256") != expected[path].get("sha256"):
            changes.append({"path": path, "change": "content_changed"})
    return changes


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--json-output", type=Path)
    parser.add_argument("--baseline", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    if not root.is_dir():
        parser.error("--root must be an existing directory")
    baseline = json.loads(args.baseline.read_text()) if args.baseline else None
    paths = inventory(root)
    if not paths:
        parser.error("No workbook or CSV files found")
    results = []
    for number, path in enumerate(paths, start=1):
        item = audit_file(path, root)
        results.append(item)
        if item["status"] != "pass" or number % 25 == 0 or number == len(paths):
            print(f"[{number}/{len(paths)}] {item['status']}: {item['path']}", flush=True)
    differences = compare_baseline(results, baseline) if baseline else []
    failures = [item for item in results if item["status"] == "fail"]
    warning_count = sum(bool(item.get("warnings")) for item in results)
    report = {
        "schema_version": 1,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "All workbook and CSV files; excludes Git metadata and dependency/cache directories",
        "limitations": "Readability and structural integrity only; does not validate scientific claims or prove original completeness without source comparison",
        "summary": {"files": len(results), "csv": sum(item["kind"] == "csv" for item in results),
                    "workbooks": sum(item["kind"] != "csv" for item in results),
                    "passed": len(results) - len(failures), "failed": len(failures),
                    "files_with_warnings": warning_count, "baseline_differences": len(differences)},
        "files": results,
        "baseline_differences": differences,
    }
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report["summary"], sort_keys=True))
    for item in failures:
        print(f"FAIL {item['path']}: {item['error']}", file=sys.stderr)
    for item in results:
        for warning in item.get("warnings", []):
            print(f"NOTE {item['path']}: {warning}", file=sys.stderr)
    if differences:
        print(f"Baseline differs for {len(differences)} file(s)", file=sys.stderr)
    return 1 if failures or differences else 0


if __name__ == "__main__":
    raise SystemExit(main())
